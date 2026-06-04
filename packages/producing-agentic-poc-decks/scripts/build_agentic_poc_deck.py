#!/usr/bin/env python3
"""Build an editable Farnsworth-style agentic POC deck from a JSON brief.

This script is intentionally self-contained so it can serve as both a runnable
template and an implementation reference for future deck work.

Usage:
    python scripts/build_agentic_poc_deck.py \
      --brief assets/example-brief.json \
      --out output/agentic-poc-deck.pptx
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, Iterable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


FONT = "Helvetica Neue"
BLACK = RGBColor(0x0A, 0x0A, 0x0A)
BODY = RGBColor(0x1C, 0x1C, 0x1C)
MUTED = RGBColor(0x53, 0x53, 0x53)
GRID = RGBColor(0xD8, 0xD8, 0xD8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF6, 0xF6, 0xF6)
TEAL = RGBColor(0x1E, 0x6F, 0x68)
NAVY = RGBColor(0x17, 0x20, 0x2C)
PALE_BLUE = RGBColor(0xE8, 0xEE, 0xF6)
PALE_GREEN = RGBColor(0xEA, 0xF2, 0xEF)
PALE_TAN = RGBColor(0xF4, 0xE5, 0xD9)
ORANGE = RGBColor(0xB9, 0x6B, 0x33)
RED = RGBColor(0xF0, 0x60, 0x60)

SLIDE_W = 13.333
SLIDE_H = 7.5
LEFT = 0.72
RIGHT = 0.72
TOP = 0.38
CONTENT_TOP = 0.97


def inch(value: float):
    return Inches(value)


def set_alpha(line, val: str = "18000") -> None:
    """Set line alpha using DrawingML. 18000 means roughly 18% opacity."""

    ln = line._element.spPr.get_or_add_ln()
    solid = ln.find(qn("a:solidFill"))
    if solid is None:
        solid = OxmlElement("a:solidFill")
        ln.insert(0, solid)
    srgb = solid.find(qn("a:srgbClr"))
    if srgb is None:
        srgb = OxmlElement("a:srgbClr")
        solid.append(srgb)
    srgb.set("val", "D8D8D8")
    for child in list(srgb):
        if child.tag == qn("a:alpha"):
            srgb.remove(child)
    alpha = OxmlElement("a:alpha")
    alpha.set("val", val)
    srgb.append(alpha)


def add_grid(slide) -> None:
    width = SLIDE_W
    height = SLIDE_H
    for i in range(1, 16):
        x = width * i / 16
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, inch(x), inch(height / 9), inch(x), inch(height)
        )
        line.name = "Opulent Hairline Grid Line"
        line.line.color.rgb = GRID
        line.line.width = Pt(0.25)
        set_alpha(line)
    for i in range(1, 9):
        y = height * i / 9
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, inch(0), inch(y), inch(width), inch(y)
        )
        line.name = "Opulent Hairline Grid Line"
        line.line.color.rgb = GRID
        line.line.width = Pt(0.25)
        set_alpha(line)


def textbox(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    size: float,
    *,
    bold: bool = False,
    color: RGBColor = BODY,
    align=PP_ALIGN.LEFT,
    name: str | None = None,
):
    shape = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    for idx, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    return shape


def rect(slide, x: float, y: float, w: float, h: float, fill=WHITE, line=BLACK):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, inch(x), inch(y), inch(w), inch(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(0.8)
    return shape


def rule(slide, x: float, y: float, w: float, color=BLACK, h: float = 0.045):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, inch(x), inch(y), inch(w), inch(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def header(slide, label: str, page: int, brief: dict[str, Any]) -> None:
    textbox(slide, LEFT, TOP, 7.9, 0.18, label.upper(), 8.5, bold=True, color=BLACK)
    marker = f"{brief.get('confidential_label', 'CONFIDENTIAL / OPULENT x CLIENT')} / {page:02d}"
    textbox(slide, 6.3, TOP, 5.9, 0.18, marker, 8, color=MUTED, align=PP_ALIGN.RIGHT)


def source(slide, text: str) -> None:
    if text:
        textbox(slide, LEFT, 7.17, 11.5, 0.16, text, 6.5, color=MUTED)


def bullets(slide, items: Iterable[str], x: float, y: float, w: float, gap: float = 0.54):
    for idx, item in enumerate(items):
        textbox(slide, x, y + idx * gap, 0.25, 0.22, "-", 14, color=BLACK)
        textbox(slide, x + 0.35, y + idx * gap, w - 0.35, 0.34, item, 11.5, color=BODY)


def add_blank(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_grid(slide)
    return slide


def slide_cover(prs, brief, page):
    slide = add_blank(prs)
    header(slide, "", page, brief)
    textbox(slide, LEFT, 1.26, 7.9, 1.55, brief["title"], 42, bold=True, color=BLACK)
    textbox(slide, LEFT, 3.23, 7.4, 0.75, brief["subtitle"], 15.5, color=BODY)
    rule(slide, LEFT, 4.6, 2.1)
    textbox(slide, LEFT, 4.9, 5.1, 0.2, f"Prepared by {brief.get('partner', 'Opulent')} for {brief.get('client', 'Client')}  -  {brief.get('date', '')}", 8.5, bold=True, color=BLACK)
    rect(slide, 8.65, 1.15, 3.55, 3.95)
    textbox(slide, 8.95, 1.55, 2.9, 0.2, "CORE SIGNAL", 8.5, bold=True, color=BLACK)
    textbox(slide, 8.95, 1.9, 2.9, 0.95, "Survey quality becomes\na trusted, repeatable\nworkflow.", 19, bold=True, color=BLACK)
    textbox(slide, 8.95, 3.5, 2.85, 0.85, "Review every respondent, surface only the exceptions, and attach evidence to every call - so PMs decide faster and stand behind the result.", 10.5)
    nums = brief.get("validated_numbers", {})
    stats = [("0", "recommended exclusions"), (nums.get("respondents", "1,036"), "sample respondents"), ("3", "delivery options")]
    for i, (num, label) in enumerate(stats):
        x = 5.85 + i * 2.55
        textbox(slide, x, 5.62, 1.6, 0.45, num, 26, bold=True, color=BLACK, align=PP_ALIGN.CENTER)
        textbox(slide, x, 6.27, 1.6, 0.18, label, 7.5, color=MUTED, align=PP_ALIGN.CENTER)


def slide_exec(prs, brief, page):
    slide = add_blank(prs)
    header(slide, "Executive Summary", page, brief)
    textbox(slide, LEFT, CONTENT_TOP, 10.7, 0.45, "Delivering a trusted survey-quality workflow.", 26, bold=True, color=BLACK)
    items = [
        "The current GPT flags poor-quality respondents, but PMs still re-review the file because open-end judgment is not trusted enough for company-wide adoption.",
        "The model is not the differentiator - the calibrated, auditable workflow around it is.",
        "PMs should review only respondents that truly need judgment, each with a reason, confidence level, and recommended action.",
        "The goal is to cut PM labor, preserve or improve quality, and keep bad respondents out of client deliverables.",
    ]
    bullets(slide, items, LEFT, 2.25, 6.9, 0.72)
    rect(slide, 8.0, 2.0, 3.7, 3.6)
    textbox(slide, 8.3, 2.35, 3.0, 0.2, "THE GOAL", 8.5, bold=True, color=BLACK)
    textbox(slide, 8.3, 2.7, 3.0, 0.95, "Exception review,\nnot full-file review.", 20, bold=True, color=BLACK)
    textbox(slide, 8.3, 4.4, 3.0, 0.8, "The target is a system that reviews everyone, flags only the exceptions, and gets sharper from PM overrides.", 10.5)


def slide_challenge(prs, brief, page):
    slide = add_blank(prs)
    header(slide, "The Challenge Today", page, brief)
    textbox(slide, LEFT, CONTENT_TOP, 11.0, 0.55, "Today's cleaning runs through Decipher, ChatGPT, and a full PM review.", 25, bold=True, color=BLACK)
    steps = [
        ("1", "Field", "Surveys are fielded in Decipher during active fielding."),
        ("2", "Download", "PMs download the respondent data file from Decipher."),
        ("3", "Upload", "The file is uploaded to ChatGPT for quality screening."),
        ("4", "Annotate", "ChatGPT returns an annotated version flagging issues."),
        ("5", "Review", "PMs review the whole file to confirm nothing was missed."),
    ]
    for i, (num, title, body) in enumerate(steps):
        x = LEFT + i * 2.18
        textbox(slide, x, 2.95, 1.7, 0.18, num, 10, bold=True, color=MUTED)
        textbox(slide, x, 3.25, 1.7, 0.3, title, 15, bold=True, color=BLACK)
        textbox(slide, x, 3.75, 1.75, 0.6, body, 9.5)
    textbox(slide, LEFT, 5.6, 2.1, 0.18, "THE BOTTLENECK", 8.5, bold=True, color=BLACK)
    textbox(slide, 3.3, 5.55, 8.6, 0.55, brief["customer_words"]["bottleneck"], 11.5)
    source(slide, "Source: customer-provided process notes/email; validate exact baseline during Discovery.")


def slide_sample(prs, brief, page):
    slide = add_blank(prs)
    header(slide, "Sample File Review", page, brief)
    nums = brief.get("validated_numbers", {})
    textbox(slide, LEFT, CONTENT_TOP, 10.9, 0.5, f"Your sample file: {nums.get('respondents', '1,036')} respondents and {nums.get('recommended_exclusions', '0')} recommended exclusions.", 24, bold=True, color=BLACK)
    textbox(slide, LEFT, 2.0, 10.5, 0.45, "The current workflow can surface signals, but it stops short of trusted decision support.", 13)
    metrics = [
        (nums.get("respondents", "1,036"), "REVIEWED"),
        ("988", "KEEP"),
        ("45", "LIGHT REVIEW"),
        ("3", "REVIEW CLOSELY"),
        (nums.get("recommended_exclusions", "0"), "DELETE / EXCLUDE"),
    ]
    for i, (value, label) in enumerate(metrics):
        x = LEFT + i * 2.2
        textbox(slide, x, 3.25, 1.65, 0.45, value, 28, bold=True, color=BLACK, align=PP_ALIGN.CENTER)
        textbox(slide, x, 4.2, 1.65, 0.16, label, 7, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
        if i > 0:
            rule(slide, x - 0.25, 3.42, 0.01, color=MUTED, h=1.1)
    textbox(slide, LEFT, 5.45, 2.3, 0.18, "THE ADOPTION GAP", 8.5, bold=True, color=BLACK)
    textbox(slide, 3.25, 5.35, 8.4, 0.65, f"The AI flagged {nums.get('open_end_concern', '785')} respondents on open-ends - {nums.get('moderate_plus_open_end', '557')} at moderate-or-higher concern - yet recommended zero exclusions. PMs are left to reconcile 'look harder, cut none' by hand.", 11.2)
    source(slide, "Source: sample file analysis; keep/remove decisions and flag counts should be recalculated for each new client.")


def slide_fraud(prs, brief, page):
    slide = add_blank(prs)
    header(slide, "Why Bad Data Gets In", page, brief)
    textbox(slide, LEFT, CONTENT_TOP, 10.9, 0.68, "Why bad data gets in - and why a static GPT cannot keep it out.", 25, bold=True, color=BLACK)
    rect(slide, LEFT, 1.95, 5.55, 3.8)
    rect(slide, 6.95, 1.95, 5.55, 3.8)
    textbox(slide, LEFT + 0.3, 2.3, 4.9, 0.2, "THE FRAUD PROBLEM", 13, bold=True, color=BLACK)
    textbox(slide, LEFT + 0.3, 2.65, 4.8, 2.25, "Bad actors take surveys to earn the incentive, not answer honestly.\nThey impersonate the target profile.\nThey speed, straightline, and duplicate entries.\nThey increasingly use AI to generate plausible open-ends.", 10.3)
    textbox(slide, LEFT + 0.3, 5.0, 4.75, 0.45, "This is where fraud is hardest to judge and where trust breaks down.", 10, bold=True, color=BLACK)
    textbox(slide, 7.25, 2.3, 4.9, 0.2, "STATIC GPT ANNOTATES; AGENTS DECIDE", 12.5, bold=True, color=BLACK)
    textbox(slide, 7.25, 2.65, 4.8, 2.45, "Today's GPT uses heuristic open-end tells that AI-assisted fraud can defeat. It recommends little or nothing, so PMs re-read everything.\n\nThe agentic approach reviews every respondent across all signals, cross-checks claimed identity, attaches quoted evidence and confidence, routes exceptions to PMs, and learns from overrides.", 9.8)
    rect(slide, LEFT, 6.05, 11.8, 0.55, fill=LIGHT)
    textbox(slide, LEFT + 0.3, 6.18, 11.0, 0.2, "OUTCOME: Catch more real fraud, trust the call, and review only what needs a human.", 12.5, bold=True, color=BLACK)


def slide_expect(prs, brief, page):
    slide = add_blank(prs)
    header(slide, "What To Expect", page, brief)
    textbox(slide, LEFT, CONTENT_TOP, 10.9, 0.62, "PM review shifts from re-reading the full file to a focused pass on respondents that need judgment.", 23, bold=True, color=BLACK)
    items = [
        "PMs review far fewer respondents because low-risk cases auto-clear.",
        "Open-end review becomes evidence-backed: quoted text, reason, and confidence.",
        "Exclusion decisions become more consistent across PMs.",
        "Each run produces a defensible audit trail.",
        "Because cleaning runs a few times a week, per-cycle overhead drops every week, not once.",
    ]
    bullets(slide, items, LEFT, 2.25, 6.9, 0.62)
    rect(slide, 8.0, 2.25, 3.7, 3.3)
    textbox(slide, 8.3, 2.6, 3.0, 0.2, "SUCCESS TARGET", 8.5, bold=True, color=BLACK)
    textbox(slide, 8.3, 3.0, 3.0, 0.3, "Interim batches", 13.5, bold=True, color=BLACK)
    textbox(slide, 8.3, 3.35, 3.0, 0.6, "Move PM review toward a focused exception pass; set the reduction target in Discovery.", 9.5)
    textbox(slide, 8.3, 4.55, 3.0, 0.3, "Final closeout", 13.5, bold=True, color=BLACK)
    textbox(slide, 8.3, 4.9, 3.0, 0.45, "Clean data, exclusion log, and summary PMs can approve with minimal rework.", 9.5)


def slide_success(prs, brief, page):
    slide = add_blank(prs)
    header(slide, "How We Measure Success", page, brief)
    textbox(slide, LEFT, CONTENT_TOP, 10.9, 0.45, "Four dimensions define a successful pilot.", 25, bold=True, color=BLACK)
    cards = [
        ("Labor reduction", "Target: review-time reduction set in Discovery. Phase 1 establishes the real baseline."),
        ("Quality & trust", "AI/PM agreement by flag type, override rate, false negatives, and false positives."),
        ("Adoption", "Company-wide adoption - the goal the current GPT never reached because it was too slow and not trusted."),
        ("Client defensibility", "Data quality as a differentiator the customer can stand behind with clients."),
    ]
    for i, (title, body) in enumerate(cards):
        x = LEFT + (i % 2) * 5.6
        y = 2.1 + (i // 2) * 2.0
        rect(slide, x, y, 5.05, 1.55)
        rule(slide, x + 0.28, y + 0.3, 0.55)
        textbox(slide, x + 0.28, y + 0.55, 4.4, 0.25, title, 15, bold=True, color=BLACK)
        textbox(slide, x + 0.28, y + 0.95, 4.45, 0.45, body, 9.5)


def slide_technical(prs, brief, page):
    slide = add_blank(prs)
    header(slide, "Technical Overview / Survey Fraud Agents", page, brief)
    textbox(slide, LEFT, CONTENT_TOP, 8.4, 0.75, "Construct agents that convert subtle survey fraud signals into evidence-backed exception review.", 24, bold=True, color=BLACK)
    textbox(slide, LEFT, 1.75, 8.2, 0.35, "Objective: protect deliverables by catching patterns humans miss at scale while moving PMs from full-file re-review to a prioritized queue.", 10.5, color=MUTED)
    rect(slide, LEFT, 2.55, 2.1, 2.9)
    textbox(slide, LEFT + 0.2, 2.85, 1.7, 0.2, "INPUTS", 7.5, bold=True, color=TEAL)
    textbox(slide, LEFT + 0.2, 3.1, 1.7, 0.45, "Survey evidence stream", 13, bold=True, color=BLACK)
    textbox(slide, LEFT + 0.2, 3.85, 1.65, 0.8, "Decipher export + datamap\nLikert batteries\nOpen-end verbatims\nTiming, device, IP, source", 8.2, color=MUTED)
    rect(slide, 3.55, 2.35, 5.5, 3.5)
    textbox(slide, 3.85, 2.6, 4.8, 0.2, "SPECIALIST FRAUD-AGENT ENSEMBLE", 7.5, bold=True, color=TEAL)
    textbox(slide, 3.85, 2.9, 4.7, 0.35, "Each agent produces a normalized suspicion score and evidence fields for PM review.", 8.8, color=MUTED)
    boxes = [
        (3.85, 3.35, "Response Pattern", "Straightlining, acquiescence, effort."),
        (6.3, 3.35, "Longitudinal / Panel", "Contradictions, velocity, fatigue."),
        (3.85, 4.15, "Demographic Plausibility", "Profile vs. behavior checks."),
        (6.3, 4.15, "Verbatim Integrity", "Near-duplicates, AI-like text."),
        (3.85, 4.95, "Real-Time Calibration", "Dynamic thresholds by study type."),
    ]
    for x, y, title, body in boxes:
        w = 2.25 if title != "Real-Time Calibration" else 4.7
        fill = PALE_GREEN if "Pattern" in title or "Integrity" in title else PALE_BLUE if "Panel" in title or "Real" in title else PALE_TAN
        rect(slide, x, y, w, 0.58, fill=fill, line=fill)
        textbox(slide, x + 0.15, y + 0.14, w - 0.25, 0.18, title, 10.5, bold=True, color=BLACK)
        textbox(slide, x + 0.15, y + 0.38, w - 0.25, 0.14, body, 6.6, color=MUTED)
    rect(slide, 9.45, 2.55, 2.95, 2.9, fill=NAVY, line=NAVY)
    textbox(slide, 9.72, 2.88, 2.35, 0.2, "DECISIONING LAYER", 7.5, bold=True, color=RGBColor(0x76, 0xD0, 0xC4))
    textbox(slide, 9.72, 3.18, 2.3, 0.55, "Risk fusion + action routing", 17, bold=True, color=WHITE)
    textbox(slide, 9.72, 4.0, 2.3, 0.5, "Combine agent scores with calibrated thresholds by study and PM override history.", 8.5, color=RGBColor(0xD7, 0xDE, 0xE8))
    for i, (c, t) in enumerate([(RGBColor(0x76, 0xD0, 0xC4), "Auto-keep: low-risk completes"), (RGBColor(0xF0, 0xA0, 0x5A), "Review: ambiguous cases"), (RED, "Exclude: high-confidence fraud")]):
        rect(slide, 9.72, 4.75 + i * 0.32, 0.07, 0.07, fill=c, line=c)
        textbox(slide, 9.9, 4.7 + i * 0.32, 2.1, 0.15, t, 7.5, color=WHITE)
    rect(slide, 3.55, 6.05, 5.5, 0.72)
    textbox(slide, 3.85, 6.25, 1.8, 0.15, "EVIDENCE CONTRACT", 7.5, bold=True, color=TEAL)
    textbox(slide, 3.85, 6.48, 4.8, 0.16, "For every surfaced respondent: version, triggering answers, cluster, contradiction, confidence band, action.", 7.5, color=MUTED)
    source(slide, "Business alignment: agents reduce missed bad respondents, explain exceptions, and protect legitimate respondents.")


def slide_how_help(prs, brief, page):
    slide = add_blank(prs)
    header(slide, "How We'd Help", page, brief)
    textbox(slide, LEFT, CONTENT_TOP, 10.9, 0.55, "A workflow we build and run for you - Decipher in, review-ready outputs back.", 24, bold=True, color=BLACK)
    steps = [
        ("01", "Ingest", "Read export + datamap -> so no PM setup is needed."),
        ("02", "Review", "Rules + open-end AI review everyone -> so quality is checked, not sampled."),
        ("03", "Recommend", "Reason + confidence + evidence -> so decisions are defensible."),
        ("04", "Route", "Route only exceptions to PMs -> so they review judgment cases."),
        ("05", "Export", "Clean data + log + audit summary -> so closeout is sign-off."),
        ("06", "Improve", "PM overrides tune thresholds -> so each run sharpens."),
    ]
    for i, (num, title, body) in enumerate(steps):
        x = LEFT + i * 1.92
        textbox(slide, x, 3.0, 1.4, 0.18, num, 10, bold=True, color=MUTED)
        textbox(slide, x, 3.35, 1.4, 0.24, title, 14.5, bold=True, color=BLACK)
        textbox(slide, x, 3.8, 1.45, 0.72, body, 8.5)
    textbox(slide, LEFT, 5.85, 1.4, 0.18, "PRINCIPLE", 8.5, bold=True, color=BLACK)
    textbox(slide, 2.3, 5.8, 9.6, 0.35, "Light for the customer to adopt because the infrastructure runs on our side.", 11.5)


def slide_opportunity(prs, brief, page):
    slide = add_blank(prs)
    header(slide, "The Opportunity", page, brief)
    textbox(slide, LEFT, CONTENT_TOP, 10.9, 0.5, "Move PMs from full-file review to exception review.", 25, bold=True, color=BLACK)
    rect(slide, LEFT, 2.1, 5.3, 3.2)
    rect(slide, 6.7, 2.1, 5.3, 3.2)
    textbox(slide, LEFT + 0.3, 2.45, 4.7, 0.2, "TODAY", 8.5, bold=True, color=BLACK)
    textbox(slide, LEFT + 0.3, 2.85, 4.7, 0.25, "Full-file review", 15, bold=True, color=BLACK)
    bullets(slide, ["AI annotates the file.", "PMs re-read everything.", "Trust burden sits with PMs.", "Logs assembled by hand."], LEFT + 0.3, 3.35, 4.5, 0.43)
    textbox(slide, 7.0, 2.45, 4.7, 0.2, "TARGET STATE", 8.5, bold=True, color=BLACK)
    textbox(slide, 7.0, 2.85, 4.7, 0.25, "Exception review", 15, bold=True, color=BLACK)
    bullets(slide, ["Agents cross-check identity, behavior, and open-ends.", "PMs review only exceptions.", "Evidence + confidence on each call.", "Overrides feed the loop."], 7.0, 3.35, 4.5, 0.43)
    rect(slide, LEFT, 5.85, 11.3, 0.45, fill=LIGHT)
    textbox(slide, LEFT + 0.25, 5.97, 10.7, 0.16, "LOOP: PM overrides feed the system so each run gets sharper and cheaper than the last.", 10.5, bold=True, color=BLACK)


def slide_options(prs, brief, page):
    slide = add_blank(prs)
    header(slide, "Implementation Options", page, brief)
    textbox(slide, LEFT, CONTENT_TOP, 10.9, 0.5, "Three ways to deliver it - from lightest to most scalable.", 25, bold=True, color=BLACK)
    cards = [
        ("01", "Email Report", "You send the export; we return an email summary with counts, top issues, and high-priority respondents.", "Outcome: fastest proof, zero change management."),
        ("02", "Enhanced Excel Package", "A workbook with a front-tab summary and exception queue. PMs approve or override directly in the file.", "Outcome: labor reduction now, in tools PMs already use."),
        ("03", "Lightweight Review App", "A review app showing project status, evidence, and approve/override controls.", "Outcome: company-wide rollout + override data."),
    ]
    for i, (num, title, body, outcome) in enumerate(cards):
        x = LEFT + i * 3.85
        rect(slide, x, 2.2, 3.35, 3.2)
        textbox(slide, x + 0.28, 2.55, 2.8, 0.18, num, 10, bold=True, color=MUTED)
        textbox(slide, x + 0.28, 2.9, 2.8, 0.28, title, 15.5, bold=True, color=BLACK)
        textbox(slide, x + 0.28, 3.75, 2.75, 0.75, body, 9.2)
        textbox(slide, x + 0.28, 4.85, 2.75, 0.35, outcome, 8.8, bold=True, color=BLACK)
    textbox(slide, LEFT, 5.95, 2.5, 0.18, "RECOMMENDED START", 8.5, bold=True, color=BLACK)
    textbox(slide, 3.8, 5.85, 8.2, 0.5, f"Recommended start: {brief['pilot'].get('recommended_start', 'Enhanced Excel + email summary')}. The payout has three parts we will size with you: incentive leakage, PM labor, and client-trust risk.", 10.5)


def slide_engagement(prs, brief, page):
    slide = add_blank(prs)
    header(slide, "The Engagement", page, brief)
    textbox(slide, LEFT, CONTENT_TOP, 10.9, 0.6, "A focused, paid pilot - prove value on real exports, then scale.", 25, bold=True, color=BLACK)
    phases = [
        ("01", "Phase 1", "Discovery & Baseline POC", "1 WEEK", "Confirm workflow, flag taxonomy, review baseline, and success metrics."),
        ("02", "Phase 2", "Workflow Integration", "1-2 WEEKS", "Implement review process, evidence-backed recommendations, and PM-ready outputs."),
        ("03", "Phase 3", "Calibration & Impact", "1-2 WEEKS", "Compare recommendations to PM decisions, measure impact, and recommend rollout path."),
    ]
    for i, (num, title, sub, dur, body) in enumerate(phases):
        x = LEFT + i * 4.0
        rule(slide, x, 2.55, 3.15)
        textbox(slide, x, 2.9, 0.7, 0.18, num, 9.5, bold=True, color=MUTED)
        textbox(slide, x, 3.25, 3.2, 0.28, title, 17, bold=True, color=BLACK)
        textbox(slide, x, 3.75, 3.2, 0.22, sub, 11, bold=True, color=BLACK)
        textbox(slide, x, 4.1, 3.2, 0.18, dur, 8.5, bold=True, color=MUTED)
        textbox(slide, x, 4.55, 3.15, 0.75, body, 10)
    textbox(slide, LEFT, 6.35, 1.5, 0.18, "OBJECTIVE", 8.5, bold=True, color=BLACK)
    textbox(slide, 2.4, 6.3, 9.2, 0.25, "Reduce manual review time while preserving or improving respondent-quality standards.", 11.5)


def slide_discovery(prs, brief, page):
    slide = add_blank(prs)
    header(slide, "Discovery - To Align On", page, brief)
    textbox(slide, LEFT, CONTENT_TOP, 10.9, 0.5, "Questions we will work through together before drafting the SOW.", 25, bold=True, color=BLACK)
    cards = [
        ("Pain & current workflow", "How long does PM review take today?\nHow many cleaning cycles in a typical week?\nHow many PMs for adoption?"),
        ("Inputs & outputs", "Which export formats should we support?\nIs the cycle interim batch, full closeout, or both?\nWhich output should the pilot target first?"),
        ("Decision logic & calibration", "Where does the keep/remove decision live?\nCan you share historical decisions?\nWhat confidence is required for exclusion?"),
        ("Success metrics", "What review-time reduction makes the pilot a win?\nWhat agreement level builds trust?\nWhat audit output helps you stand behind quality?"),
    ]
    for i, (title, body) in enumerate(cards):
        x = LEFT + (i % 2) * 5.6
        y = 2.2 + (i // 2) * 2.05
        rect(slide, x, y, 5.05, 1.65)
        textbox(slide, x + 0.3, y + 0.32, 4.4, 0.24, title, 14, bold=True, color=BLACK)
        textbox(slide, x + 0.3, y + 0.75, 4.45, 0.7, body, 9.2)


def slide_addons(prs, brief, page):
    slide = add_blank(prs)
    header(slide, "Potential Add-ons", page, brief)
    textbox(slide, LEFT, CONTENT_TOP, 10.9, 0.48, "Where this can go after the pilot proves value.", 25, bold=True, color=BLACK)
    addons = [
        ("01", "Pre-field intake review", "Better screeners before fielding."),
        ("02", "Respondent friction", "Reduce straightlining, speeding, and fatigue."),
        ("03", "Decipher integration", "Scheduled interim review during fielding."),
        ("04", "Client-facing QA", "Make quality assurance a sales advantage."),
        ("05", "Automated report QA", "Cut unsupported claims before clients see them."),
    ]
    for i, (num, title, body) in enumerate(addons):
        x = LEFT + i * 2.25
        rule(slide, x, 2.65, 1.8)
        textbox(slide, x, 2.95, 0.5, 0.18, num, 8.5, bold=True, color=MUTED)
        textbox(slide, x, 3.3, 1.85, 0.45, title, 11.3, bold=True, color=BLACK)
        textbox(slide, x, 4.1, 1.85, 0.55, body, 8.5)
    textbox(slide, LEFT, 6.35, 11.0, 0.22, "Extend the workflow upstream into intake quality and downstream into reporting only after the pilot earns it.", 10, color=MUTED)


def slide_commit(prs, brief, page):
    slide = add_blank(prs)
    header(slide, "Commitment & Next Steps", page, brief)
    textbox(slide, LEFT, CONTENT_TOP, 10.9, 0.55, "What we need from you to launch the pilot.", 25, bold=True, color=BLACK)
    items = [
        "Pursue a focused pilot around PM review reduction.",
        "Provide representative exports and datamaps.",
        "Share final keep/remove decisions or help build a gold-standard set.",
        "Identify one to three PMs to test the workflow.",
        "Agree on pilot success metrics before the build begins.",
        "Stay open to phased rollout if the pilot proves out.",
    ]
    bullets(slide, items, LEFT, 2.2, 6.8, 0.62)
    rect(slide, 8.25, 2.3, 3.6, 3.85)
    textbox(slide, 8.55, 2.7, 3.0, 0.18, "THE OUTCOME", 8.5, bold=True, color=BLACK)
    textbox(slide, 8.55, 3.1, 3.0, 1.8, "Internally: less operational effort and a workflow every PM adopts.\n\nExternally: defensible data quality clients can trust.", 15, bold=True, color=BLACK)


def slide_thanks(prs, brief, page):
    slide = add_blank(prs)
    header(slide, "Thank You", page, brief)
    textbox(slide, LEFT, 2.15, 7.5, 0.7, "Thank you.", 46, bold=True, color=BLACK)
    textbox(slide, LEFT, 3.35, 7.8, 0.95, f"We appreciate the opportunity to help {brief.get('client', 'your team')} make quality more trusted, defensible, and easier to adopt.", 21, color=BODY)
    rule(slide, LEFT, 4.65, 2.1)
    textbox(slide, LEFT, 5.15, 5.0, 0.22, f"{brief.get('partner', 'Opulent')} x {brief.get('client', 'Client')}", 12, bold=True, color=BLACK)


SLIDE_BUILDERS = [
    slide_cover,
    slide_exec,
    slide_challenge,
    slide_sample,
    slide_fraud,
    slide_expect,
    slide_success,
    slide_technical,
    slide_how_help,
    slide_opportunity,
    slide_options,
    slide_engagement,
    slide_discovery,
    slide_addons,
    slide_commit,
    slide_thanks,
]


def build_deck(brief: dict[str, Any], out: Path) -> None:
    prs = Presentation()
    prs.slide_width = inch(SLIDE_W)
    prs.slide_height = inch(SLIDE_H)
    for idx, builder in enumerate(SLIDE_BUILDERS, 1):
        if builder is slide_technical and not brief.get("slides", {}).get("include_technical_overview", True):
            continue
        if builder is slide_addons and not brief.get("slides", {}).get("include_addons", True):
            continue
        if builder is slide_thanks and not brief.get("slides", {}).get("include_thank_you", True):
            continue
        builder(prs, brief, len(prs.slides) + 1)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--brief", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    args = parser.parse_args()
    brief = json.loads(args.brief.read_text())
    build_deck(brief, args.out)
    print(args.out)


if __name__ == "__main__":
    main()
